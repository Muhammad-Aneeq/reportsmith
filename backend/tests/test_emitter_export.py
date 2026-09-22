"""The emitter's export shape — the same layout the intake template will accept.

One format for both doors, so that `generate → export → upload → parse` round-trips.
The parser lands in P3; what can be asserted now is that the emitted grid is complete,
ordered, platform-stable and losslessly re-readable as Decimals. P3 adds the other half
of the round-trip.
"""

from __future__ import annotations

import csv
import io
from decimal import Decimal

import pytest

from ledgerfab.statements import STATEMENT_LINES, emit_statements
from ledgerfab.statements.export import to_csv, to_rows, to_xlsx


@pytest.fixture(scope="module")
def statements():  # type: ignore[no-untyped-def]
    return emit_statements("squeeze", seed=42, periods=4, grain="quarter")


def test_header_is_the_three_keys_plus_one_column_per_period(statements) -> None:  # type: ignore[no-untyped-def]
    header = to_rows(statements)[0]
    assert header[:3] == ["statement", "line_code", "label"]
    assert header[3:] == list(statements.period_ids)


def test_every_declared_line_is_exported_exactly_once(statements) -> None:  # type: ignore[no-untyped-def]
    rows = to_rows(statements)[1:]
    exported = [(row[0], row[1]) for row in rows]
    assert len(exported) == len(set(exported)), "a line was exported twice"
    assert {code for _, code in exported} == {ln.code for ln in STATEMENT_LINES}


def test_rows_follow_statement_declaration_order(statements) -> None:  # type: ignore[no-untyped-def]
    """So a human reading the file sees a P&L that runs revenue → net income, rather
    than an alphabetical list in which `cogs` precedes `revenue`."""
    codes = [row[1] for row in to_rows(statements)[1:]]
    declared = [ln.code for ln in STATEMENT_LINES]
    assert codes == declared


def test_amounts_round_trip_to_the_same_decimals(statements) -> None:  # type: ignore[no-untyped-def]
    """The point of the export: nothing is lost on the way out."""
    reader = csv.reader(io.StringIO(to_csv(statements)))
    header = next(reader)
    period_ids = header[3:]
    for row in reader:
        for period_id, cell in zip(period_ids, row[3:], strict=True):
            assert Decimal(cell) == statements.amount(period_id, row[1])


def test_csv_line_endings_are_platform_stable(statements) -> None:  # type: ignore[no-untyped-def]
    """A Windows checkout and a Linux CI runner must produce identical bytes, or
    file-hash comparisons start failing for reasons unrelated to the data."""
    text = to_csv(statements)
    assert "\r" not in text
    assert text.endswith("\n")


def test_amounts_carry_no_presentation_chrome(statements) -> None:  # type: ignore[no-untyped-def]
    """No thousands separators, no currency symbol, no parenthesised negatives.

    Those are a *parsing* problem and the P3 parser owns them. Emitting them here would
    make the round-trip test quietly depend on the parser's leniency rather than on the
    two formats agreeing.
    """
    for row in to_rows(statements)[1:]:
        for cell in row[3:]:
            assert not any(ch in cell for ch in ",()£$ "), cell
            assert cell.count(".") == 1 and len(cell.split(".")[1]) == 2


def test_xlsx_is_a_readable_workbook(statements) -> None:  # type: ignore[no-untyped-def]
    from openpyxl import load_workbook

    workbook = load_workbook(io.BytesIO(to_xlsx(statements)))
    sheet = workbook["statements"]
    assert sheet.max_row == len(to_rows(statements))
    assert [cell.value for cell in sheet[1]][:3] == ["statement", "line_code", "label"]
    # Revenue is the first declared line, so row 2 of the sheet.
    assert sheet.cell(row=2, column=2).value == "revenue"
    assert isinstance(sheet.cell(row=2, column=4).value, (int, float))
